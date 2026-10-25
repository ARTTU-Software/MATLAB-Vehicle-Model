"""Offline fixture tests; no MATLAB required or accessed."""
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
import zipfile

spec = importlib.util.spec_from_file_location('inventory', Path(__file__).with_name('inspect_saved_project.py'))
mod = importlib.util.module_from_spec(spec)
spec.loader.exec_module(mod)


class SavedInventoryTests(unittest.TestCase):
    def test_literals_do_not_evaluate_expressions(self):
        self.assertEqual(mod.literal('[1 2;3 4]')['value'], [[1, 2], [3, 4]])
        self.assertEqual(mod.literal("system('touch unwanted')")['status'], 'expression_unresolved')
        self.assertTrue(mod.literal('[1 2 3]', 2)['value_omitted'])
        self.assertNotIn('value', mod.literal('[1 2 3]', 2))

    def test_slx_hierarchy_mask_lines_and_opaque_data(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root/'example.slx'
            with zipfile.ZipFile(path, 'w') as z:
                z.writestr('simulink/systems/system_root.xml', '''<System>
                  <Block Name="A/B" SID="1" BlockType="SubSystem">
                    <Mask><MaskParameter Name="K"><Value>Vehicle.mass</Value></MaskParameter></Mask>
                    <System Ref="system_1"/>
                  </Block>
                  <Block Name="Out" SID="3" BlockType="Outport"/>
                  <Line><P Name="Src">1#out:1</P><Branch><P Name="Dst">3#in:1</P></Branch></Line>
                </System>''')
                z.writestr('simulink/systems/system_1.xml', '''<System>
                  <Block Name="Gain" SID="2" BlockType="Gain"><P Name="Gain">2</P></Block>
                </System>''')
                z.writestr('simulink/blockDiagram.json', '{"BlockDiagram":{"InitFcn":"do_not_execute"}}')
                z.writestr('simulink/modelworkspace.mxarray', b'opaque')
            result = mod.Inspector(root, root/'out', [], 100).slx(path)
            self.assertEqual(len(result['blocks']), 3)
            self.assertEqual(result['blocks'][1]['path'], 'example/A//B/Gain')
            self.assertEqual(result['blocks'][1]['parameters']['Gain'], '2')
            self.assertEqual(len(result['connections']), 1)
            self.assertEqual(result['coverage']['visited_system_members'], 2)
            self.assertEqual(result['embedded_data'][0]['status'], 'opaque_matlab_serialization')

    def test_legacy_mdl_inline_system(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); path = root/'legacy.mdl'
            path.write_text('Model { Name "legacy" System { Block { BlockType Gain Name "Gain" Gain "5" } } }')
            result = mod.Inspector(root, root/'out', [], 100).mdl(path)
            self.assertEqual(result['blocks'][0]['path'], 'legacy/Gain')
            self.assertEqual(result['blocks'][0]['parameters']['Gain'], '5')

    def test_saved_mat_nested_struct_and_large_array(self):
        import numpy as np
        from scipy.io import savemat
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); path = root/'data.mat'
            savemat(path, {'Vehicle': {'mass': 320.0, 'curve': np.arange(150)}, 'big': np.arange(300)})
            inspector = mod.Inspector(root, root/'out', [], 100)
            inspector.mat_data(path, 'data.mat')
            mass = next(v for v in inspector.variables if v['target'] == 'Vehicle.mass')
            curve = next(v for v in inspector.variables if v['target'] == 'Vehicle.curve')
            self.assertEqual(mass['value'], [[320.0]])
            self.assertTrue(curve['value_omitted'])
            self.assertFalse(inspector.issues)

    def test_excel_formula_preserved_without_execution(self):
        import openpyxl
        from openpyxl.worksheet.formula import ArrayFormula
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); path = root/'sm_car_ARTTU_Test.xlsx'
            wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Part'
            ws.append(['Type', None, 'Test'])
            ws.cell(5,1,'Mass'); ws.cell(5,3,320); ws.cell(5,4,'kg')
            ws.cell(6,1,'Derived'); ws['C6'] = ArrayFormula(ref='C6:C7', text='=C5*2')
            wb.save(path); wb.close()
            inspector = mod.Inspector(root, root/'out', [], 100)
            inspector.workbook(path)
            encoded = json.dumps(mod.safe_json(inspector.variables))
            self.assertIn('ArrayFormula', encoded)
            self.assertIn('=C5*2', encoded)


if __name__ == '__main__':
    unittest.main()
