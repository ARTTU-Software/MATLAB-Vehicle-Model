#!/usr/bin/env python3
"""Offline inventory of saved Simulink projects. Never starts or calls MATLAB.

Requires numpy, scipy, openpyxl and h5py for data-file inspection. SLX parsing
uses observed saved-file formats, not a supported replacement for Simulink APIs.
Uninterpreted/opaque records are explicitly reported rather than evaluated.
"""
from __future__ import annotations

import argparse
import csv
import gzip
import hashlib
import io
import json
import math
import os
from pathlib import Path
import re
import sys
from datetime import datetime, timezone
import xml.etree.ElementTree as ET
import zipfile

SCHEMA = 1
EXCLUDE = {'.git', 'work', 'slprj', 'spinfo', '__pycache__', 'SimulationStudies'}
MODEL_EXT = {'.slx', '.mdl', '.slxp'}
SOURCE_EXT = {'.m', '.ssc', '.mlx'}
DATA_EXT = {'.mat', '.xlsx', '.xlsm', '.sldd', '.slddc', '.p'}
NUM = r'[+-]?(?:(?:\d+(?:\.\d*)?|\.\d+)(?:[eEdD][+-]?\d+)?|Inf|NaN)'


def digest(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b''):
            h.update(chunk)
    return h.hexdigest()


def safe_json(value):
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    if isinstance(value, complex):
        return {'real': safe_json(value.real), 'imag': safe_json(value.imag)}
    if isinstance(value, dict):
        return {str(k): safe_json(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [safe_json(v) for v in value]
    if isinstance(value, (datetime,)):
        return value.isoformat()
    if isinstance(value, bytes):
        return {'bytes_hex': value.hex()}
    if value.__class__.__module__.startswith('openpyxl'):
        return {'excel_object_type': type(value).__name__,
                'attributes': safe_json(vars(value))}
    if hasattr(value, 'item'):
        return safe_json(value.item())
    return value


def write_json(path, obj):
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + '.tmp')
    content = (json.dumps(safe_json(obj), ensure_ascii=False, indent=2, allow_nan=False) + '\n').encode()
    tmp.write_bytes(gzip.compress(content, mtime=0) if path.suffix == '.gz' else content)
    tmp.replace(path)


def write_lines(path, rows):
    tmp = path.with_suffix(path.suffix + '.tmp')
    opener = gzip.open if path.suffix == '.gz' else open
    with opener(tmp, 'wt', encoding='utf-8') as f:
        for row in rows:
            f.write(json.dumps(safe_json(row), ensure_ascii=False, allow_nan=False) + '\n')
    tmp.replace(path)


def read_json(path):
    data = path.read_bytes()
    return json.loads(gzip.decompress(data) if path.suffix == '.gz' else data)


def literal(text, limit=100):
    """Parse only simple literals, never MATLAB expressions or function calls."""
    if not isinstance(text, str):
        return {'status': 'stored_value', 'value': safe_json(text)}
    s = text.strip().rstrip(';').strip()
    if s in ('true', 'false'):
        return {'status': 'literal', 'value': s == 'true', 'shape': [1, 1]}
    if re.fullmatch(NUM, s, re.I):
        return {'status': 'literal', 'value': safe_json(float(re.sub('[dD]', 'e', s))), 'shape': [1, 1]}
    if re.fullmatch(r"'(?:[^']|'')*'", s) or re.fullmatch(r'"(?:[^"]|"")*"', s):
        v = s[1:-1].replace(s[0] * 2, s[0])
        return {'status': 'literal', **({'value': v} if len(v) <= limit else {'length': len(v), 'value_omitted': True})}
    if s.startswith('[') and s.endswith(']'):
        rows = []
        for row in s[1:-1].split(';'):
            tokens = row.replace(',', ' ').split()
            if any(not re.fullmatch(NUM, t, re.I) for t in tokens):
                break
            rows.append([safe_json(float(re.sub('[dD]', 'e', t))) for t in tokens])
        else:
            sizes = {len(r) for r in rows}
            if len(sizes) == 1:
                count = sum(len(r) for r in rows)
                return {'status': 'literal', 'shape': [len(rows), len(rows[0])], 'elements': count,
                        **({'value': rows} if count <= limit else {'value_omitted': True})}
    return {'status': 'expression_unresolved'}


def properties(el):
    return {p.get('Name', ''): p.text or '' for p in el.findall('P')}


def xml_record(el):
    return {'tag': el.tag, 'attributes': dict(el.attrib),
            **({'text': el.text} if el.text and el.text.strip() else {}),
            **({'children': [xml_record(c) for c in el]} if len(el) else {})}


def escaped(name):
    return name.replace('/', '//')


def paths_under(root, extensions, excluded=()):
    if not root.exists():
        return []
    found = []
    for folder, dirs, files in os.walk(root):
        dirs[:] = sorted(d for d in dirs if d not in excluded and not d.startswith('.'))
        for name in sorted(files):
            if name.startswith('~$'):
                continue  # Office lock files are not saved workbooks.
            p = Path(folder) / name
            if p.suffix.lower() in extensions:
                found.append(p)
    return sorted(found)


class Inspector:
    def __init__(self, root, out, library_roots, limit):
        self.root, self.out, self.limit = root, out, limit
        self.library_roots = library_roots
        self.files = paths_under(root, MODEL_EXT | SOURCE_EXT | DATA_EXT, EXCLUDE)
        self.catalog = {}
        for p in [x for x in self.files if x.suffix.lower() in MODEL_EXT]:
            self.catalog.setdefault(p.stem, []).append(p)
        for folder in library_roots:
            for p in paths_under(folder, MODEL_EXT):
                self.catalog.setdefault(p.stem, []).append(p)
        self.models, self.variables, self.sources, self.workbooks = [], [], [], []
        self.dependencies, self.issues, self.fingerprints = [], [], {}
        self.model_by_source = {}
        self.cached_models, self.cached_fingerprints, self.cached_embedded = {}, {}, []
        self.reused_models = 0
        self.previous_model_records = []
        manifest = out/'manifest.json'
        if manifest.exists():
            old = json.loads(manifest.read_text())
            self.previous_model_records = json.loads((out/'models.json').read_text())
            if old.get('inspector_sha256') == digest(__file__) and old.get('small_value_limit') == limit:
                self.cached_fingerprints = json.loads((out/'fingerprints.json').read_text())
                self.cached_models = {m['source']: m for m in json.loads((out/'models.json').read_text())}
                with gzip.open(out/'variables.jsonl.gz', 'rt', encoding='utf-8') as f:
                    self.cached_embedded = [v for line in f if '::' in line and (v := json.loads(line)).get('kind','').startswith(('mat_', 'hdf5_'))]

    def source(self, path):
        try:
            return path.relative_to(self.root).as_posix()
        except ValueError:
            for i, root in enumerate(self.library_roots):
                try:
                    return f'@library{i}/' + path.relative_to(root).as_posix()
                except ValueError:
                    pass
            return str(path)

    def fingerprint(self, path):
        source = self.source(path)
        if source not in self.fingerprints:
            self.fingerprints[source] = {'sha256': digest(path), 'bytes': path.stat().st_size}
        return self.fingerprints[source]

    def issue(self, source, kind, detail):
        self.issues.append({'source': source, 'kind': kind, 'detail': str(detail)})

    def block(self, b, path, member):
        p = properties(b)
        result = {'path': path, 'name': b.get('Name'), 'sid': b.get('SID'),
                  'type': b.get('BlockType'), 'source_member': member, 'parameters': p,
                  'stored_literals': {k: v for k, t in p.items() if (v := literal(t, self.limit))['status'] == 'literal'},
                  'details': [xml_record(c) for c in b if c.tag not in ('P', 'System')],
                  'variant_status': 'saved_structure_not_compiled'}
        return result

    def slx(self, path):
        source = self.source(path)
        r = {'source': source, 'name': path.stem, 'blocks': [], 'connections': [],
             'annotations': [], 'metadata': {}, 'algorithms': [], 'embedded_data': [], 'gaps': []}
        with zipfile.ZipFile(path) as z:
            names = z.namelist()
            systems = {}
            for name in names:
                if name.startswith('simulink/systems/') and name.endswith('.xml') and '/_rels/' not in name:
                    systems[Path(name).stem] = (ET.fromstring(z.read(name)), name)
            visited = set()

            def walk(system, prefix, member, stack=()):
                visited.add(member)
                for b in system.findall('Block'):
                    bp = prefix + '/' + escaped(b.get('Name', '?'))
                    r['blocks'].append(self.block(b, bp, member))
                    for sub in b.findall('System'):
                        ref = sub.get('Ref')
                        if ref:
                            if ref in stack:
                                r['gaps'].append({'kind': 'cyclic_system_ref', 'ref': ref, 'path': bp})
                            elif ref in systems:
                                element, loc = systems[ref]
                                walk(element, bp, loc, (*stack, ref))
                            else:
                                r['gaps'].append({'kind': 'missing_system_ref', 'ref': ref, 'path': bp})
                        else:
                            walk(sub, bp, member, stack)
                for line in system.findall('Line'):
                    r['connections'].append({'system': prefix, 'source_member': member, 'record': xml_record(line)})
                for ann in system.findall('Annotation'):
                    r['annotations'].append({'system': prefix, 'sid': ann.get('SID'), 'properties': properties(ann)})

            if 'system_root' in systems:
                walk(*[systems['system_root'][0], path.stem, systems['system_root'][1]], ('system_root',))
            elif 'simulink/blockdiagram.xml' in names:
                old = ET.fromstring(z.read('simulink/blockdiagram.xml'))
                for s in old.findall('.//Model/System') + old.findall('.//Library/System') + old.findall('.//Subsystem/System'):
                    walk(s, path.stem, 'simulink/blockdiagram.xml')
            for key, (s, loc) in systems.items():
                if loc not in visited:
                    r['gaps'].append({'kind': 'unattached_system', 'source_member': loc})
                    walk(s, path.stem + '/@unattached/' + key, loc)
            for name in names:
                if name.endswith('.mxarray'):
                    r['embedded_data'].append({'member': name, 'bytes': z.getinfo(name).file_size,
                                               'status': 'opaque_matlab_serialization'})
                elif name.endswith('.mat'):
                    self.mat_data(io.BytesIO(z.read(name)), source + '::' + name)
                elif name.startswith('simulink/stateflow/') and name.endswith('.xml'):
                    r['algorithms'].append({'source_member': name, 'record': xml_record(ET.fromstring(z.read(name)))})
                elif name.startswith('simulink/') and name.count('/') == 1 and name.endswith(('.json', '.xml')):
                    data = z.read(name)
                    if name.endswith('.json'):
                        r['metadata'][name] = json.loads(data)
                    else:
                        el = ET.fromstring(data)
                        # Systems already recorded with full paths; avoid duplication.
                        for parent in el.iter():
                            for child in list(parent):
                                if child.tag == 'System' and (child.find('Block') is not None):
                                    parent.remove(child)
                        r['metadata'][name] = xml_record(el)
            r['coverage'] = {'saved_block_count': len(r['blocks']), 'system_members': len(systems),
                             'visited_system_members': len(visited), 'opaque_members': len(r['embedded_data'])}
        return r

    def mdl(self, path):
        """Read legacy brace-based MDL structure without evaluating its contents."""
        text = path.read_text(errors='replace')
        tokens = re.findall(r'"(?:\\.|[^"\\])*"|[{}]|[^\s{}"]+', text)
        pos = 0
        def value(t):
            if t.startswith('"'):
                try:
                    return json.loads(t)
                except ValueError:
                    return t[1:-1]
            return t
        def parse():
            nonlocal pos
            children, props = [], {}
            while pos < len(tokens):
                key = tokens[pos]; pos += 1
                if key == '}':
                    break
                if pos >= len(tokens):
                    break
                token = tokens[pos]; pos += 1
                if token == '{':
                    child = parse(); child['tag'] = key; children.append(child)
                else:
                    v = value(token)
                    while pos < len(tokens) and tokens[pos].startswith('"'):
                        v += value(tokens[pos]); pos += 1
                    props[key] = v
            return {'properties': props, 'children': children}
        tree = parse()
        r = {'source': self.source(path), 'name': path.stem, 'blocks': [], 'connections': [],
             'annotations': [], 'metadata': {'legacy_mdl': tree}, 'algorithms': [], 'embedded_data': [], 'gaps': []}
        def walk(node, prefix):
            for child in node['children']:
                p = child['properties']; tag = child['tag']
                if tag == 'Block':
                    bp = prefix + '/' + escaped(p.get('Name', '?'))
                    r['blocks'].append({'path': bp, 'name': p.get('Name'), 'sid': p.get('SID'),
                        'type': p.get('BlockType'), 'parameters': p, 'source_member': 'legacy_mdl',
                        'details': [c for c in child['children'] if c['tag'] != 'System'],
                        'stored_literals': {k: v for k, t in p.items() if (v := literal(t, self.limit))['status'] == 'literal'},
                        'variant_status': 'saved_structure_not_compiled'})
                    walk(child, bp)
                elif tag == 'Line':
                    r['connections'].append({'system': prefix, 'record': child})
                else:
                    walk(child, prefix)
        walk(tree, path.stem)
        r['coverage'] = {'saved_block_count': len(r['blocks']), 'legacy_mdl_parser': True}
        return r

    def inspect_models(self):
        queue = [p for p in self.files if p.suffix.lower() in MODEL_EXT]
        seen = set()
        while queue:
            p = queue.pop(0)
            if p in seen:
                continue
            seen.add(p)
            source = self.source(p)
            self.fingerprint(p)
            try:
                if p.suffix.lower() == '.slxp':
                    self.issue(source, 'protected_model', 'Interior unavailable to offline inspection')
                    continue
                cached = self.cached_models.get(source)
                if cached and self.cached_fingerprints.get(source) == self.fingerprints[source]:
                    r = read_json(self.out/cached['record'])
                    self.variables.extend(v for v in self.cached_embedded if v['source'].startswith(source+'::'))
                    self.reused_models += 1
                else:
                    r = self.slx(p) if p.suffix.lower() == '.slx' else self.mdl(p)
                self.model_by_source[source] = r
                self.models.append(r)
                for b in r['blocks']:
                    params = b['parameters']
                    for key in ('SourceBlock', 'ModelName', 'ReferencedSubsystem'):
                        ref = params.get(key)
                        if not ref or ref in ('<not specified>', '<none>'):
                            continue
                        rootname = ref.split('/')[0] if key == 'SourceBlock' else Path(ref).stem
                        if ref.startswith('<'):
                            self.dependencies.append({'source':source,'block':b['path'],'kind':key,'reference':ref,
                                'candidates':[], 'status':'template_placeholder'})
                            continue
                        candidates = [p] if rootname == '$bdroot' else self.catalog.get(rootname, [])
                        local = [c for c in candidates if c.is_relative_to(self.root)]
                        chosen = local if local else candidates
                        dep = {'source': source, 'block': b['path'], 'kind': key, 'reference': ref,
                               'candidates': [self.source(c) for c in chosen],
                               'status': 'unique_file_candidate_path_not_runtime_verified' if len(chosen) == 1 else ('ambiguous' if chosen else 'missing')}
                        self.dependencies.append(dep)
                        queue.extend(chosen)
                for gap in r['gaps']:
                    self.issue(source, gap['kind'], gap)
                if r['embedded_data']:
                    self.issue(source, 'opaque_embedded_data', r['embedded_data'])
            except Exception as exc:
                self.issue(source, 'model_read_error', repr(exc))
            if len(seen) % 25 == 0:
                print(f'Inspected {len(seen)} saved model/library files', flush=True)

    def source_text(self, source, text):
        """Conservative assignment index, not an evaluator or full MATLAB parser."""
        lines = text.splitlines(); scope = 'script'; buffer = ''; start = 0; block_comment = False
        record = {'source': source, 'functions': [], 'assignment_count': 0,
                  'status': 'lexical_source_only', 'dynamic_operations': [], 'signal_accesses': []}
        for line_no, line in enumerate(lines, 1):
            stripped = line.strip()
            if stripped == '%{': block_comment = True
            if block_comment:
                if stripped == '%}': block_comment = False
                continue
            if not stripped or stripped.startswith('%'): continue
            # Remove comments only outside strings; transpose quotes are preserved.
            clean = ''; quote = None; i = 0
            while i < len(line):
                c = line[i]
                if quote:
                    clean += c
                    if c == quote:
                        if i + 1 < len(line) and line[i + 1] == quote:
                            clean += quote; i += 1
                        else: quote = None
                elif c == '%': break
                else:
                    clean += c
                    if c == '"' or (c == "'" and (i == 0 or line[i-1].isspace() or line[i-1] in '=([{,;')):
                        quote = c
                i += 1
            clean = clean.strip()
            if not buffer: start = line_no
            buffer += ' ' + clean
            if buffer.rstrip().endswith('...'):
                buffer = buffer.rstrip()[:-3]; continue
            statement = buffer.strip(); buffer = ''
            if statement.startswith('function '):
                record['functions'].append({'line': start, 'signature': statement})
                scope = statement; continue
            if re.search(r'\b(eval|evalin|assignin|load|run|set_param)\s*\(', statement):
                record['dynamic_operations'].append({'line': start, 'statement': statement})
            if re.search(r'logsout|\.Values\.|\.Data\b|\.Time\b', statement):
                record['signal_accesses'].append({'line': start, 'statement': statement})
            match = re.match(r'^([A-Za-z]\w*(?:\.[A-Za-z]\w*|\([^=]*?\)|\{[^=]*?\})*|\[[^\]]+\])\s*=(?!=)\s*(.*)$', statement)
            if match:
                lhs, rhs = match.groups()
                self.variables.append({'source': source, 'line': start, 'scope': scope, 'target': lhs,
                    'kind': 'source_assignment', 'expression': rhs, 'literal': literal(rhs, self.limit),
                    'dependencies_candidates': sorted(set(re.findall(r'\b[A-Za-z]\w*(?:\.[A-Za-z]\w*)+', rhs))),
                    'verification': 'source_statement_not_effective_workspace_value'})
                record['assignment_count'] += 1
        self.sources.append(record)

    def mat_data(self, data, source):
        import numpy as np
        from scipy.io import whosmat, loadmat
        def flatten(value, target, ancestors=()):
            if type(value).__name__ == 'MatlabOpaque':
                self.variables.append({'source': source, 'target': target, 'kind': 'mat_opaque_object',
                    'shape': list(value.shape), 'class_metadata': str(value['_Class'].tolist()) if '_Class' in value.dtype.names else None,
                    'verification': 'opaque_matlab_object_not_deserialized'})
                self.issue(source, 'opaque_matlab_object', target)
                return
            if id(value) in ancestors:
                self.issue(source, 'cyclic_mat_value', target); return
            if hasattr(value, '_fieldnames'):
                self.variables.append({'source': source, 'target': target, 'kind': 'mat_structure', 'fields': value._fieldnames})
                for name in value._fieldnames:
                    flatten(getattr(value, name), target + '.' + name, (*ancestors, id(value)))
            elif isinstance(value, np.ndarray) and (value.dtype.names or value.dtype == object):
                self.variables.append({'source': source, 'target': target, 'kind': 'mat_container', 'shape': list(value.shape), 'dtype': str(value.dtype)})
                for idx, item in enumerate(value.flat):
                    suffix = '' if value.size == 1 else f'{{linear_C_{idx}}}'
                    if value.dtype.names:
                        for field in value.dtype.names:
                            flatten(item[field], target + suffix + '.' + field, (*ancestors, id(value)))
                    else:
                        flatten(item, target + suffix, (*ancestors, id(value)))
            else:
                arr = np.asarray(value)
                entry = {'source': source, 'target': target, 'kind': 'mat_value', 'shape': list(arr.shape),
                         'dtype': str(arr.dtype), 'elements': int(arr.size), 'verification': 'saved_data_not_base_workspace'}
                if arr.dtype.kind in 'biufcUS' and arr.size <= self.limit:
                    entry['value'] = safe_json(arr.tolist())
                else: entry['value_omitted'] = True
                self.variables.append(entry)
        try:
            try:
                headers = whosmat(data)
            except TypeError:
                # SciPy cannot list some MCOS objects, but can expose saved names/classes.
                if hasattr(data, 'seek'): data.seek(0)
                saved = loadmat(data, struct_as_record=False, squeeze_me=False)
                for name, value in saved.items():
                    if not name.startswith('__'): flatten(value, name)
                return
            for name, shape, cls in headers:
                self.variables.append({'source': source, 'target': name, 'kind': 'mat_variable', 'shape': list(shape), 'class': cls})
                # Numeric arrays need not be loaded if their values will be omitted.
                if cls not in ('struct', 'cell', 'object') and math.prod(shape) > self.limit:
                    continue
                if hasattr(data, 'seek'): data.seek(0)
                value = loadmat(data, variable_names=[name], struct_as_record=False, squeeze_me=False)[name]
                flatten(value, name)
        except NotImplementedError:
            self.hdf_mat(data, source)
        except Exception as exc:
            self.issue(source, 'mat_read_error', repr(exc))

    def hdf_mat(self, data, source):
        import h5py
        import numpy as np
        try:
            with h5py.File(data, 'r') as f:
                def visit(name, obj):
                    row = {'source': source, 'target': '/' + name, 'kind': 'hdf5_mat_storage',
                           'attributes': {k: str(v) for k, v in obj.attrs.items()},
                           'verification': 'storage_path_not_matlab_field_mapping'}
                    if isinstance(obj, h5py.Dataset):
                        row.update(shape=list(obj.shape), dtype=str(obj.dtype), elements=int(obj.size))
                        if obj.size <= self.limit:
                            v = obj[()]
                            if h5py.check_dtype(ref=obj.dtype):
                                row['references'] = [f[r].name if r else None for r in np.asarray(v).flat]
                            elif obj.dtype.kind in 'biufc': row['value'] = safe_json(np.asarray(v).tolist())
                            else: row['value_omitted'] = True
                        else: row['value_omitted'] = True
                    self.variables.append(row)
                f.visititems(visit)
        except Exception as exc:
            self.issue(source, 'hdf5_read_error', repr(exc))

    def workbook(self, path):
        import openpyxl
        source = self.source(path)
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            record = {'source': source, 'sheets': [], 'defined_names': [str(x) for x in wb.defined_names.values()]}
            for ws in wb:
                nonempty, small = 0, []
                rows = []
                for row_number, row in enumerate(ws.iter_rows(), 1):
                    values = [c.value for c in row]
                    for cell in row:
                        if cell.value is not None:
                            nonempty += 1
                            if len(small) < self.limit:
                                small.append({'cell': cell.coordinate, 'value': str(cell.value) if cell.data_type == 'f' else cell.value})
                    if path.name.startswith('sm_car_ARTTU_') and any(x is not None for x in values[:5]):
                        rows.append({'row': row_number, 'cells_A_to_E': values[:5]})
                sheet = {'name': ws.title, 'rows': ws.max_row, 'columns': ws.max_column,
                         'nonempty_cells': nonempty, 'verification': 'stored_cells_formulas_not_evaluated'}
                if nonempty <= self.limit: sheet['cells'] = small
                else: sheet['values_omitted'] = True
                if rows:
                    sheet['import_rows'] = rows
                    category = None
                    for row in rows:
                        vals = row['cells_A_to_E'] + [None]*5
                        a,b,c,unit,comment = vals[:5]
                        if row['row'] < 5 or row['row'] > 100: continue
                        if a is not None: category = str(a)
                        if category and c is not None:
                            self.variables.append({'source': source, 'sheet': ws.title, 'cell': f'C{row["row"]}',
                                'kind': 'excel_import_candidate', 'target': f'{ws.title}.{category}' + (f'.{b}' if b else '') + '.Value',
                                'workbook_category': path.stem.removeprefix('sm_car_ARTTU_'),
                                'unit': unit, 'comment': comment, 'literal': literal(c,self.limit), 'stored_cell': c,
                                'verification': 'candidate_import_field_before_assembly; SHEET/TABLE references require importer tracing'})
                record['sheets'].append(sheet)
            self.workbooks.append(record)
        finally: wb.close()

    def data_sources(self):
        for p in self.files:
            if p.suffix.lower() in MODEL_EXT: continue
            source = self.source(p); self.fingerprint(p)
            try:
                suffix = p.suffix.lower()
                if suffix in ('.m', '.ssc'):
                    self.source_text(source, p.read_text(errors='replace'))
                elif suffix == '.mlx':
                    with zipfile.ZipFile(p) as z:
                        for name in z.namelist():
                            if name.endswith('document.xml'):
                                root = ET.fromstring(z.read(name))
                                self.source_text(source + '::' + name, '\n'.join(root.itertext()))
                elif suffix in ('.xlsx', '.xlsm'): self.workbook(p)
                elif suffix == '.mat': self.mat_data(p, source)
                else: self.issue(source, 'opaque_data_or_code', 'Saved bytes fingerprinted; proprietary content not decoded')
            except Exception as exc:
                self.issue(source, 'source_read_error', repr(exc))

    def save(self):
        model_index, parameters, signals = [], [], []
        for r in self.models:
            name = hashlib.sha256(r['source'].encode()).hexdigest()[:16] + '.json.gz'
            write_json(self.out/'models'/name, r)
            model_index.append({'source': r['source'], 'name': r['name'], 'record': 'models/'+name, 'blocks': len(r['blocks']), 'coverage': r['coverage']})
            for b in r['blocks']:
                for key, value in b['parameters'].items():
                    if key in {'Position','ZOrder','FontSize','FontName','ForegroundColor','BackgroundColor'}: continue
                    parameters.append({'source': r['source'], 'block': b['path'], 'parameter': key,
                        'expression': value, 'literal': literal(value, self.limit),
                        'field_references': sorted(set(re.findall(r'\b[A-Za-z]\w*(?:\.[A-Za-z]\w*)+', value)))})
                for detail in b['details']:
                    if detail.get('tag') in ('Mask','PortProperties','PortCounts'):
                        if detail['tag'] == 'Mask':
                            for c in detail.get('children',[]):
                                if c['tag'] == 'MaskParameter':
                                    parameters.append({'source':r['source'],'block':b['path'],'mask_parameter':c['attributes'].get('Name'),'record':c})
                        else: signals.append({'source':r['source'],'block':b['path'],'record':detail})
        write_json(self.out/'models.json', model_index)
        write_json(self.out/'dependencies.json', self.dependencies)
        write_json(self.out/'sources.json', self.sources)
        write_json(self.out/'workbooks.json', self.workbooks)
        write_lines(self.out/'variables.jsonl.gz', self.variables)
        write_lines(self.out/'parameters.jsonl.gz', parameters)
        write_lines(self.out/'signals.jsonl.gz', signals)
        write_json(self.out/'gaps.json', self.issues)
        write_json(self.out/'fingerprints.json', self.fingerprints)
        # A flat view lets the MATLAB-only study scripts check freshness.
        with (self.out/'fingerprints.csv').open('w', newline='') as stream:
            writer = csv.writer(stream)
            writer.writerow(['Path', 'External', 'SHA256'])
            for source, record in self.fingerprints.items():
                external = source.startswith('@library')
                path = source
                if external:
                    alias, relative = source.split('/', 1)
                    path = str(self.library_roots[int(alias.removeprefix('@library'))] / relative)
                writer.writerow([path, int(external), record['sha256']])
        summary = {'schema_version': SCHEMA, 'generated_at_utc': datetime.now(timezone.utc).isoformat(),
            'inspection': 'saved_files_only_no_matlab', 'project_root': str(self.root),
            'library_roots': {f'@library{i}':str(p) for i,p in enumerate(self.library_roots)},
            'small_value_limit': self.limit, 'inspector_sha256': digest(__file__),
            'model_records_reused_by_fingerprint': self.reused_models,
            'excluded_directories': sorted(EXCLUDE),
            'counts': {'models_and_libraries':len(self.models), 'saved_blocks':sum(len(r['blocks']) for r in self.models),
                'connections':sum(len(r['connections']) for r in self.models), 'variable_records':len(self.variables),
                'parameter_records':len(parameters), 'source_files':len(self.sources), 'workbooks':len(self.workbooks),
                'references':len(self.dependencies), 'missing_references':sum(d['status']=='missing' for d in self.dependencies),
                'ambiguous_references':sum(d['status']=='ambiguous' for d in self.dependencies), 'gaps':len(self.issues)},
            'limitations': ['No base workspace initialized or inspected; source assignments are not runtime values.',
                'No compiled variant resolution, mask evaluation, solver execution, or signal availability validation.',
                'Offline SLX/MDL parsing does not expose undocumented defaults or opaque MATLAB serialized objects.',
                'References point to definition records; library instances are not flattened or evaluated.',
                'Lexical source indexing does not prove execution order, control flow, dynamic fields, or complete variable coverage.',
                'MAT files may contain old results; saved contents are not assumed to be current model inputs.',
                'Path precedence and library forwarding are not resolved using MATLAB; candidates remain explicit.']}
        write_json(self.out/'manifest.json', summary)
        text = ['# Saved-file inventory', '', f'Generated: {summary["generated_at_utc"]}', '',
            'Offline inspection only. MATLAB was not accessed; no project code was executed.', '',
            '## Coverage', '']
        text += [f'- {k.replace("_"," ")}: **{v:,}**' for k,v in summary['counts'].items()]
        text += ['', '## Searchable records', '',
            '- `models.json`: model/library catalog; `models/*.json.gz`: compressed JSON hierarchy, block parameters, masks, connections, callbacks, algorithms and metadata.',
            '- `parameters.jsonl.gz`: block/mask parameter lookup and candidate workspace references.',
            '- `variables.jsonl.gz`: source assignments, saved MAT structures/values, spreadsheet field candidates.',
            '- `sources.json`: function signatures, dynamic operations and logging-access statements with line numbers.',
            '- `workbooks.json`: sheets, small sheet values and vehicle importer rows.',
            '- `signals.jsonl.gz`: stored port/logging metadata; runtime signal contents are unverified.',
            '- `dependencies.json`: referenced definitions, candidate files and unresolved references.',
            '- `gaps.json`: opaque records, errors and coverage limitations.',
            '- `fingerprints.json`: SHA-256 identities of inspected saved files; `fingerprints.csv` is the flat MATLAB-readable view.', '', '## Interpretation', '']
        text += ['- '+x for x in summary['limitations']]
        text += ['', 'Use `tools/query_saved_inventory.py` to search the compressed indexes without unpacking them. JSON compression preserves all records. The count covers all project model files plus recursively discovered library definitions, not the number of instantiated blocks in sm_car.', '']
        (self.out/'README.md').write_text('\n'.join(text))
        # Remove only old uncompressed artifacts from this inspector's prior catalog.
        for old in self.previous_model_records:
            record = Path(old['record'])
            if record.parent == Path('models') and re.fullmatch(r'[0-9a-f]{16}\.json', record.name):
                (self.out/record).unlink(missing_ok=True)
        for name in ('variables.jsonl', 'parameters.jsonl', 'signals.jsonl'):
            (self.out/name).unlink(missing_ok=True)
        print(json.dumps(summary['counts'], indent=2))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root', type=Path, default=Path(__file__).resolve().parents[2])
    parser.add_argument('--output', type=Path)
    parser.add_argument('--library-root', type=Path, action='append', default=[])
    parser.add_argument('--small-limit', type=int, default=100)
    args = parser.parse_args()
    root = args.root.resolve(); out = (args.output or root/'SimulationStudies/inventory/saved').resolve()
    if not out.is_relative_to(root/'SimulationStudies'):
        parser.error('Output must be inside project SimulationStudies directory')
    if args.small_limit < 1: parser.error('small-limit must be positive')
    out.mkdir(parents=True, exist_ok=True)
    inspector = Inspector(root, out, [p.resolve() for p in args.library_root], args.small_limit)
    inspector.inspect_models(); inspector.data_sources(); inspector.save()


if __name__ == '__main__':
    main()
